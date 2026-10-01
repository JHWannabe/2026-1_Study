"""
body_composition_sum.py / body_composition_L3.py 의 세그멘테이션 결과가 제대로 됐는지
눈으로 확인하기 위한 QC 시각화 스크립트.

[방법]
  z_bounds.xlsx 기준 pubis~liver 구간을 크롭 → 그 위에서
    1) TotalSegmentator "total" task (roi_subset=["vertebrae_L3"]) 로 L3 슬라이스 탐지
    2) TotalSegmentator "tissue_4_types" task 로 SAT/VAT/skeletal_muscle/IMATA 마스크 계산
  → HU 필터링으로 IMATA/NAMA/LAMA/내장지방/피하지방 마스크 확정 (body_composition_*.py와 동일 기준)
  → 구간(pubis~liver) 전체 슬라이스 각각에 대해 CT 원본과 마스크 오버레이를 나란히 그려
    환자 1명당 PPT 파일 1개로 저장 (슬라이드 순서: liver → pubis).

[출력]
  seg_qc/{SITE_EN}/{PID}.pptx  : 환자별 전체 슬라이스 세그멘테이션 오버레이
  seg_qc/{SITE_EN}_qc_log.xlsx : 환자별 seg_status 로그

[샘플링]
  SAMPLE_N 에 정수를 지정하면 site별로 무작위 SAMPLE_N명만 처리한다 (재현을 위해 SEED 고정).
"""

import os
import io
import sys
import re
import shutil
import pickle
import logging
import tempfile
import contextlib
from difflib import SequenceMatcher
from typing import cast

import numpy as np
import pandas as pd
import pydicom
import nibabel as nib
from nibabel.nifti1 import Nifti1Image
import SimpleITK as sitk
from tqdm import tqdm

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch

from pptx import Presentation
from pptx.util import Inches

matplotlib.rcParams["font.family"] = "Malgun Gothic"
matplotlib.rcParams["axes.unicode_minus"] = False

os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"
from totalsegmentator.python_api import totalsegmentator
import torch

torch.backends.cudnn.benchmark        = True
torch.backends.cuda.matmul.allow_tf32 = True
torch.backends.cudnn.allow_tf32       = True

for _lg in ("nnunetv2", "totalsegmentator", "batchgeneratorsv2",
            "acvl_utils", "dynamic_network_architectures"):
    logging.getLogger(_lg).setLevel(logging.ERROR)

DEVICE = "gpu" if torch.cuda.is_available() else "cpu"

# ── 설정 ──────────────────────────────────────────────────────────────────────

SITES = ["강남", "신촌"]
SITE_EN_MAP = {"강남": "Gangnam", "신촌": "Sinchon"}

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE_OUT_DIR = rf"{PROJECT_ROOT}\results\seg_qc"
CHECKPOINT_DIR = rf"{PROJECT_ROOT}\aec"

TASK        = "tissue_4_types"
VERT_TASK   = "total"        # L3 위치 탐지용 (roi_subset=["vertebrae_L3"])
NAMA_HU     = (30, 150)     # 정상감쇠근육 (Normal Attenuation Muscle Area)
LAMA_HU     = (-30, 30)     # 저감쇠근육   (Low Attenuation Muscle Area)
FAT_HU      = (-190, -30)   # 지방 조직 (subcutaneous_fat/torso_fat/intermuscular_fat 공통 적용)

WINDOW_CENTER  = 40    # CT 표시용 soft-tissue window
WINDOW_WIDTH   = 400

COLORS = {
    "피하지방": (1.00, 0.85, 0.00),   # yellow
    "내장지방": (0.90, 0.15, 0.15),   # red
    "IMATA":   (0.65, 0.15, 0.85),   # purple
    "NAMA":    (0.15, 0.75, 0.25),   # green
    "LAMA":    (0.15, 0.65, 0.95),   # cyan/blue
}
OVERLAY_ALPHA = 0.45

SAMPLE_N = 20    # site별 샘플 인원. None이면 전체 처리
SEED     = 42

BATCH_SIZE = 1   # 몇 명마다 체크포인트 저장할지


@contextlib.contextmanager
def _silence():
    old_out, old_err = sys.stdout, sys.stderr
    sys.stdout = sys.stderr = io.StringIO()
    try:
        yield
    finally:
        sys.stdout, sys.stderr = old_out, old_err


# ── DICOM 폴더/시리즈 매칭 (body_composition_*.py 와 동일 규칙) ────────────────

def _norm_series(s: str) -> str:
    s = s.lower()
    s = " ".join(s.split())
    s = re.sub(r"[_/]", "_", s)
    s = re.sub(r"\s*\(\s*", "(", s)
    s = re.sub(r"\s*\)\s*", ")", s)
    return s


def _get_folder_dicom_meta(folder: str) -> tuple[str | None, str | None]:
    try:
        for f in os.listdir(folder):
            fp = os.path.join(folder, f)
            if os.path.isfile(fp):
                try:
                    hdr = pydicom.dcmread(fp, stop_before_pixels=True)
                    series_desc  = str(hdr.SeriesDescription) if hasattr(hdr, "SeriesDescription") else None
                    manufacturer = str(hdr.ManufacturerModelName) if hasattr(hdr, "ManufacturerModelName") else None
                    return series_desc, manufacturer
                except Exception:
                    continue
    except Exception:
        pass
    return None, None


def build_folder_map(dicom_base: str) -> dict[str, str]:
    if not os.path.isdir(dicom_base):
        raise FileNotFoundError(f"DICOM 경로 없음: {dicom_base}")
    folder_map: dict[str, str] = {}
    for folder_name in os.listdir(dicom_base):
        full_path = os.path.join(dicom_base, folder_name)
        if not os.path.isdir(full_path):
            continue
        parts = folder_name.split("_")
        folder_map[folder_name if len(parts) == 1 else parts[1]] = full_path
    return folder_map


def find_series_dir(pid_str: str, folder_map: dict, expected_series: str | None,
                    expected_manufacturer: str | None, warnings: list) -> str | None:
    if pid_str not in folder_map:
        return None
    base = folder_map[pid_str]
    subfolders = [s for s in os.listdir(base) if os.path.isdir(os.path.join(base, s))]
    if not subfolders:
        return None
    if len(subfolders) == 1:
        return os.path.join(base, subfolders[0])

    matched     = None
    found_meta: list[tuple[str | None, str | None]] = []
    for sf in subfolders:
        sd, mf = _get_folder_dicom_meta(os.path.join(base, sf))
        found_meta.append((sd, mf))
        sd_ok = expected_series       is None or _norm_series(sd or "") == _norm_series(expected_series)
        mf_ok = expected_manufacturer is None or " ".join((mf or "").split()) == " ".join(str(expected_manufacturer).split())
        if sd_ok and mf_ok:
            matched = sf
            break
    if matched is None and expected_series is not None:
        norm_exp = _norm_series(expected_series)
        best_sf, best_ratio = None, 0.0
        for sf, (sd, _) in zip(subfolders, found_meta):
            ratio = max(
                SequenceMatcher(None, _norm_series(sd or ""), norm_exp).ratio(),
                SequenceMatcher(None, _norm_series(sf),       norm_exp).ratio(),
            )
            if ratio > best_ratio:
                best_ratio, best_sf = ratio, sf
        if best_sf:
            matched = best_sf
            warnings.append(f"series: 유사도 매칭 (ratio={best_ratio:.2f}) → '{best_sf}'")
    if matched is None:
        warnings.append(f"series: '{expected_series}' 매칭 실패 → 첫 번째 시리즈 사용 (읽힌 값: {found_meta})")
        matched = subfolders[0]
    return os.path.join(base, matched)


# ── 시각화 유틸 ────────────────────────────────────────────────────────────────

def _window_hu(hu_slice: np.ndarray) -> np.ndarray:
    lo = WINDOW_CENTER - WINDOW_WIDTH / 2
    hi = WINDOW_CENTER + WINDOW_WIDTH / 2
    clipped = np.clip(hu_slice, lo, hi)
    return (clipped - lo) / (hi - lo)


def _overlay_rgb(gray01: np.ndarray, mask_dict: dict[str, np.ndarray]) -> np.ndarray:
    rgb = np.stack([gray01] * 3, axis=-1).astype(np.float64)
    for name, mask in mask_dict.items():
        color = COLORS[name]
        for c in range(3):
            rgb[..., c] = np.where(mask, rgb[..., c] * (1 - OVERLAY_ALPHA) + color[c] * OVERLAY_ALPHA, rgb[..., c])
    return np.clip(rgb, 0, 1)


def _slice_img(hu_vol: np.ndarray, mdict: dict[str, np.ndarray], idx: int) -> np.ndarray:
    gray = _window_hu(hu_vol[:, :, idx])
    sl_masks = {name: m[:, :, idx] for name, m in mdict.items()}
    rgb = _overlay_rgb(gray, sl_masks)
    return np.transpose(rgb, (1, 0, 2))  # (X,Y,3) -> (Y,X,3) for imshow


def _legend_handles():
    return [Patch(color=c, label=n) for n, c in COLORS.items()]


def _slice_figure_png_bytes(pid: int, hu_vol: np.ndarray, mdict: dict[str, np.ndarray],
                            idx: int, global_k: int, is_l3: bool) -> io.BytesIO:
    fig, axes = plt.subplots(1, 2, figsize=(9, 4.2))
    gray = _window_hu(hu_vol[:, :, idx])
    axes[0].imshow(np.transpose(gray), cmap="gray", origin="lower")
    axes[0].set_title("CT (원본)", fontsize=9)
    axes[0].axis("off")
    axes[1].imshow(_slice_img(hu_vol, mdict, idx), origin="lower")
    axes[1].set_title("마스크 오버레이", fontsize=9)
    axes[1].axis("off")
    fig.legend(handles=_legend_handles(), loc="lower center", ncol=5, fontsize=7)
    tag = "  ★ L3" if is_l3 else ""
    fig.suptitle(f"PID {pid} | k={global_k} (local {idx}){tag}", fontsize=11)
    fig.tight_layout(rect=(0, 0.08, 1, 0.90))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=100)
    plt.close(fig)
    buf.seek(0)
    return buf


def build_patient_pptx(pid: int, hu_vol: np.ndarray, mdict: dict[str, np.ndarray],
                       lo_k: int, hi_k: int, l3_idx: int | None,
                       range_sums: dict, out_path: str) -> int:
    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(6.2)
    blank_layout = prs.slide_layouts[6]

    title_slide = prs.slides.add_slide(blank_layout)
    tf = title_slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(5)).text_frame
    tf.text = f"PID {pid}"
    for line in [
        f"구간: pubis(k={lo_k}) ~ liver(k={hi_k})  (n={hi_k - lo_k + 1} 슬라이스, 슬라이드는 liver→pubis 순서)",
        f"구간 합산 — IMATA {range_sums['IMATA']:.1f}  NAMA {range_sums['NAMA']:.1f}  LAMA {range_sums['LAMA']:.1f}  "
        f"내장지방 {range_sums['내장지방']:.1f}  피하지방 {range_sums['피하지방']:.1f}  (cm²)",
        f"L3 슬라이스: local k={l3_idx} (전역 k={lo_k + l3_idx})" if l3_idx is not None else "L3 슬라이스: 탐지 실패",
    ]:
        p = tf.add_paragraph()
        p.text = line

    n = hu_vol.shape[2]
    n_slides = 0
    for idx in range(n - 1, -1, -1):  # liver(높은 k) → pubis(낮은 k) 순서
        buf = _slice_figure_png_bytes(pid, hu_vol, mdict, idx, lo_k + idx, idx == l3_idx)
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(buf, Inches(0.3), Inches(0.2), width=Inches(9.4))
        n_slides += 1

    prs.save(out_path)
    return n_slides


# ── 핵심 처리 (환자 1명) ──────────────────────────────────────────────────────

def process_patient(zrow: pd.Series, folder_map: dict, tmp_dir: str, out_dir: str) -> dict:
    pid     = int(zrow["PatientID"])
    pid_str = str(pid)
    result  = {"PatientID": pid, "seg_status": "failed", "has_pptx": False, "n_slides": 0}
    warnings: list[str] = []

    dcm_dir = find_series_dir(pid_str, folder_map, zrow.get("series_description"),
                              zrow.get("manufacturer_model"), warnings)
    if dcm_dir is None:
        result["seg_status"] = "no_dicom"
        return result

    cropped_nii_path = os.path.join(tmp_dir, f"{pid}_cropped.nii.gz")
    seg_dir           = os.path.join(tmp_dir, f"{pid}_seg")
    vert_seg_dir       = os.path.join(tmp_dir, f"{pid}_vert_seg")

    try:
        reader     = sitk.ImageSeriesReader()
        series_ids = reader.GetGDCMSeriesIDs(dcm_dir)
        if not series_ids:
            result["seg_status"] = "no_series_uid"
            return result
        dicom_names = reader.GetGDCMSeriesFileNames(dcm_dir, series_ids[0])
        reader.SetFileNames(dicom_names)
        img = reader.Execute()

        n_k = img.GetSize()[2]
        if len(dicom_names) != n_k:
            result["seg_status"] = "invalid_volume_match"
            return result

        lo_k = int(min(zrow["pubis_k"], zrow["liver_upper_k"]))
        hi_k = int(max(zrow["pubis_k"], zrow["liver_upper_k"]))
        lo_k = max(0, min(lo_k, n_k - 1))
        hi_k = max(0, min(hi_k, n_k - 1))
        if hi_k < lo_k:
            result["seg_status"] = "invalid_bounds"
            return result

        size  = [img.GetSize()[0], img.GetSize()[1], hi_k - lo_k + 1]
        index = [0, 0, lo_k]
        cropped_img = sitk.RegionOfInterest(img, size, index)
        sitk.WriteImage(cropped_img, cropped_nii_path)

        # L3 위치 탐지
        os.makedirs(vert_seg_dir, exist_ok=True)
        with _silence():
            totalsegmentator(input=cropped_nii_path, output=vert_seg_dir, task=VERT_TASK,
                             roi_subset=["vertebrae_L3"], fast=False, quiet=True, device=DEVICE)
        l3_idx = None
        l3_mask_path = os.path.join(vert_seg_dir, "vertebrae_L3.nii.gz")
        if os.path.exists(l3_mask_path):
            l3_mask = cast(Nifti1Image, nib.load(l3_mask_path)).get_fdata() > 0.5
            area_per_slice = l3_mask.sum(axis=(0, 1))
            if np.any(area_per_slice > 0):
                l3_idx = int(np.argmax(area_per_slice))
        if l3_idx is None:
            warnings.append("L3 슬라이스 탐지 실패")

        # 체성분 세그멘테이션
        os.makedirs(seg_dir, exist_ok=True)
        with _silence():
            totalsegmentator(input=cropped_nii_path, output=seg_dir, task=TASK,
                             fast=False, quiet=True, device=DEVICE)

        class_files = {
            "subcutaneous_fat":  "subcutaneous_fat.nii.gz",
            "torso_fat":         "torso_fat.nii.gz",
            "skeletal_muscle":   "skeletal_muscle.nii.gz",
            "intermuscular_fat": "intermuscular_fat.nii.gz",
        }
        raw_masks = {}
        for cls, fname in class_files.items():
            p = os.path.join(seg_dir, fname)
            if not os.path.exists(p):
                warnings.append(f"{cls}: 세그멘테이션 결과 없음")
                result["seg_status"] = "tissue_seg_missing"
                return result
            raw_masks[cls] = cast(Nifti1Image, nib.load(p)).get_fdata() > 0.5

        cropped_nii = cast(Nifti1Image, nib.load(cropped_nii_path))
        hu_data     = cropped_nii.get_fdata()
        zooms = cast(tuple[float, ...], cropped_nii.header.get_zooms())
        sx, sy = zooms[0], zooms[1]
        pixel_area_cm2 = float(sx) * float(sy) / 100.0

        fat_hu_mask = (hu_data >= FAT_HU[0]) & (hu_data <= FAT_HU[1])
        muscle_mask = raw_masks["skeletal_muscle"]

        mdict = {
            "IMATA":    raw_masks["intermuscular_fat"] & fat_hu_mask,
            "내장지방": raw_masks["torso_fat"]          & fat_hu_mask,
            "피하지방": raw_masks["subcutaneous_fat"]   & fat_hu_mask,
            "NAMA":     muscle_mask & (hu_data >= NAMA_HU[0]) & (hu_data <= NAMA_HU[1]),
            "LAMA":     muscle_mask & (hu_data >= LAMA_HU[0]) & (hu_data <= LAMA_HU[1]),
        }

        range_sums = {name: round(float(m.sum()) * pixel_area_cm2, 2) for name, m in mdict.items()}
        pptx_out = os.path.join(out_dir, f"{pid}.pptx")
        n_slides = build_patient_pptx(pid, hu_data, mdict, lo_k, hi_k, l3_idx, range_sums, pptx_out)
        result["has_pptx"] = True
        result["n_slides"] = n_slides

        result["seg_status"] = "ok" if l3_idx is not None else "ok_no_l3"

    except Exception as e:
        result["seg_status"] = f"error:{type(e).__name__}:{e}"

    finally:
        for path in (cropped_nii_path, seg_dir, vert_seg_dir):
            if os.path.isfile(path):
                os.remove(path)
            elif os.path.isdir(path):
                shutil.rmtree(path, ignore_errors=True)

    if warnings:
        result["warnings"] = " | ".join(warnings)
        tqdm.write(f"  PID {pid}: " + " | ".join(warnings))

    return result


# ── 체크포인트 ────────────────────────────────────────────────────────────────

def _load_checkpoint(checkpoint_path: str) -> tuple[set[int], list[dict]]:
    if os.path.exists(checkpoint_path):
        with open(checkpoint_path, "rb") as f:
            data = pickle.load(f)
        processed, results = data["processed"], data["results"]
        error_pids = {r["PatientID"] for r in results if str(r.get("seg_status", "")).startswith("error")}
        if error_pids:
            processed = processed - error_pids
            results   = [r for r in results if r["PatientID"] not in error_pids]
            print(f"[체크포인트] error 상태 {len(error_pids)}명 재시도 대상으로 포함")
        print(f"[체크포인트] {len(processed)}명 완료 → 이어서 시작")
        return processed, results
    return set(), []


def _save_checkpoint(checkpoint_path: str, processed: set[int], results: list[dict]) -> None:
    os.makedirs(os.path.dirname(checkpoint_path), exist_ok=True)
    with open(checkpoint_path, "wb") as f:
        pickle.dump({"processed": processed, "results": results}, f)


def _write_log(results: list[dict], log_path: str) -> None:
    if not results:
        return
    df = pd.DataFrame(results).drop_duplicates(subset=["PatientID"], keep="last")
    df = df.sort_values("PatientID").reset_index(drop=True)
    os.makedirs(os.path.dirname(log_path), exist_ok=True)
    df.to_excel(log_path, sheet_name="qc_log", index=False)


# ── 사이트별 처리 ──────────────────────────────────────────────────────────────

def run_site(site: str) -> None:
    site_en = SITE_EN_MAP[site]
    dicom_base  = rf"E:\영상제공\{site}\{site}_axial"
    data_dir    = rf"C:\Users\jhjun\OneDrive\Desktop\2026-1_Study\연구코드\data\{site}"
    zbounds_path = rf"{data_dir}\aec\{site}_z_bounds.xlsx"
    merged_path  = rf"{PROJECT_ROOT}\data\{site_en}\{site_en}.xlsx"

    out_dir = os.path.join(BASE_OUT_DIR, site_en)
    os.makedirs(out_dir, exist_ok=True)
    checkpoint_path = os.path.join(CHECKPOINT_DIR, f"{site}_seg_qc_checkpoint.pkl")
    log_path = os.path.join(BASE_OUT_DIR, f"{site_en}_qc_log.xlsx")

    merged_pids = set(pd.read_excel(merged_path, sheet_name="metadata")["PatientID"].astype(int))

    zdf = pd.read_excel(zbounds_path)
    zdf = zdf[zdf["seg_status"].astype(str) == "ok"]
    zdf = zdf[zdf["PatientID"].astype(int).isin(merged_pids)].reset_index(drop=True)

    if SAMPLE_N is not None and len(zdf) > SAMPLE_N:
        zdf = zdf.sample(n=SAMPLE_N, random_state=SEED).sort_values("PatientID").reset_index(drop=True)

    print(f"[{site}] 대상: {len(zdf)}명")

    folder_map = build_folder_map(dicom_base)

    processed, results = _load_checkpoint(checkpoint_path)
    pending = zdf[~zdf["PatientID"].isin(processed)]
    print(f"[{site}] 완료: {len(processed)}명  |  처리 예정: {len(pending)}명")

    n_processed = 0
    with tempfile.TemporaryDirectory() as tmp_dir:
        with tqdm(total=len(zdf), initial=len(processed), desc=f"세그멘테이션 QC ({site})") as pbar:
            for _, zrow in pending.iterrows():
                pid = int(zrow["PatientID"])
                try:
                    result = process_patient(zrow, folder_map, tmp_dir, out_dir)
                except Exception as e:
                    tqdm.write(f"  [ERROR] PID {pid}: {type(e).__name__}: {e}")
                    result = {"PatientID": pid, "seg_status": f"error:{type(e).__name__}:{e}",
                              "has_pptx": False, "n_slides": 0}

                results.append(result)
                processed.add(pid)
                n_processed += 1
                pbar.update(1)
                tqdm.write(f"  PID {pid}: {result['seg_status']}")

                if n_processed % BATCH_SIZE == 0 or n_processed == len(pending):
                    _save_checkpoint(checkpoint_path, processed, results)
                    _write_log(results, log_path)

    _write_log(results, log_path)
    counts = pd.Series([r["seg_status"] for r in results]).value_counts().to_dict()
    print(f"\n[{site} 완료] 총 {len(results)}명  |  {counts}  |  {out_dir}")


def main():
    for site in SITES:
        run_site(site)


if __name__ == "__main__":
    main()
