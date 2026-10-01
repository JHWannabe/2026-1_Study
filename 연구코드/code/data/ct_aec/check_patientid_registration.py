"""
D:\\데이터서비스팀 요청(이홍선) 폴더명(prefix_patientID_연도_CT)에서 patientID를 추출하여,
"통합 문서.xlsx"의 모든 시트에 있는 '연구등록번호' 컬럼과 전부 일치(포함)하는지 확인한다.

- 폴더명의 patientID가 각 시트의 '연구등록번호' 값들 안에 전부 존재하면 통과.
- 존재하지 않는 patientID가 있으면 그 개수와 목록을 결과 파일로 저장한다.
"""

import os
import re
import csv
import openpyxl

BASE_DIR = r"D:\데이터서비스팀 요청(이홍선)"
EXCEL_PATH = os.path.join(BASE_DIR, "통합 문서.xlsx")
REG_COL_NAME = "연구등록번호"

SUMMARY_CSV = os.path.join(BASE_DIR, "patientID_check_summary.csv")
MISSING_DIR = os.path.join(BASE_DIR, "patientID_check_missing_by_sheet")

INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


def sanitize_filename(name):
    """시트명을 안전한 파일명으로 변환"""
    return INVALID_FILENAME_CHARS.sub("_", name).strip()


def get_folder_patient_ids(base_dir):
    """폴더명(prefix_patientID_연도_CT)에서 patientID 집합을 추출"""
    patient_ids = set()
    for name in os.listdir(base_dir):
        full_path = os.path.join(base_dir, name)
        if not os.path.isdir(full_path) or full_path == MISSING_DIR:
            continue
        parts = name.split("_")
        if len(parts) < 4:
            print(f"[경고] 폴더명 형식이 예상과 다름: {name}")
            continue
        patient_id = parts[1]
        if not patient_id.isdigit():
            print(f"[경고] patientID가 숫자가 아님: {name}")
            continue
        patient_ids.add(int(patient_id))
    return patient_ids


def get_sheet_registration_ids(ws):
    """시트의 헤더에서 '연구등록번호' 컬럼을 찾아 값 집합을 반환"""
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    if REG_COL_NAME not in header:
        return None
    col_idx = header.index(REG_COL_NAME)

    reg_ids = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        value = row[col_idx]
        if value is None:
            continue
        try:
            reg_ids.add(int(value))
        except (ValueError, TypeError):
            print(f"[경고] {ws.title} 시트의 {REG_COL_NAME} 값 변환 실패: {value!r}")
    return reg_ids


def check_sheet(label, ws, folder_patient_ids, summary_rows):
    """시트 하나에 대해 patientID 일치 여부를 확인하고 결과를 summary_rows/CSV에 반영"""
    reg_ids = get_sheet_registration_ids(ws)

    if reg_ids is None:
        print(f"[{label}] '{REG_COL_NAME}' 컬럼을 찾을 수 없음 -> 건너뜀")
        summary_rows.append([label, len(folder_patient_ids), "N/A", "N/A", "컬럼 없음"])
        return 0

    missing = sorted(folder_patient_ids - reg_ids)
    matched_count = len(folder_patient_ids) - len(missing)

    status = "일치" if not missing else "불일치"
    print(f"[{label}] 전체 {len(folder_patient_ids)} / 일치 {matched_count} / 불일치 {len(missing)}")

    summary_rows.append([label, len(folder_patient_ids), matched_count, len(missing), status])

    if missing:
        sheet_csv = os.path.join(MISSING_DIR, f"{sanitize_filename(label)}.csv")
        with open(sheet_csv, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(["missing_patientID"])
            writer.writerows([[pid] for pid in missing])

    return len(missing)


def main():
    print(f"폴더 스캔 중: {BASE_DIR}")
    folder_patient_ids = get_folder_patient_ids(BASE_DIR)
    print(f"폴더에서 추출한 고유 patientID 개수: {len(folder_patient_ids)}")

    os.makedirs(MISSING_DIR, exist_ok=True)

    summary_rows = []
    total_missing = 0

    print(f"엑셀 로드 중: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        total_missing += check_sheet(sheet_name, wb[sheet_name], folder_patient_ids, summary_rows)
    wb.close()

    with open(SUMMARY_CSV, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["sheet_name", "total_folder_patientID", "matched_count", "missing_count", "status"])
        writer.writerows(summary_rows)
    print(f"\n요약 저장 완료: {SUMMARY_CSV}")
    print(f"불일치 patientID 목록 저장 완료 (시트별 파일, 총 {total_missing}건): {MISSING_DIR}")


if __name__ == "__main__":
    main()
