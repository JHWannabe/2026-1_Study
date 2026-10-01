"""
폴더에서 추출한 patientID(13387개)를 행으로, "통합 문서.xlsx"의 각 시트를 열로 하여
해당 patientID가 그 시트의 '연구등록번호'에 존재하면 1, 없으면 0으로 표기한 매트릭스를 저장한다.
"""

import os
import openpyxl
from check_patientid_registration import (
    BASE_DIR,
    EXCEL_PATH,
    get_folder_patient_ids,
    get_sheet_registration_ids,
)

OUTPUT_XLSX = os.path.join(BASE_DIR, "patientID_sheet_matrix.xlsx")


def main():
    print(f"폴더 스캔 중: {BASE_DIR}")
    folder_patient_ids = sorted(get_folder_patient_ids(BASE_DIR))
    print(f"patientID 개수: {len(folder_patient_ids)}")

    print(f"엑셀 로드 중: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    sheet_reg_ids = {}
    for sheet_name in wb.sheetnames:
        reg_ids = get_sheet_registration_ids(wb[sheet_name])
        sheet_reg_ids[sheet_name] = reg_ids if reg_ids is not None else set()
        print(f"[{sheet_name}] 연구등록번호 고유값 {len(sheet_reg_ids[sheet_name])}개")
    wb.close()

    out_wb = openpyxl.Workbook()
    ws = out_wb.active
    ws.title = "matrix"

    sheet_names = list(sheet_reg_ids.keys())
    ws.append(["patientID"] + sheet_names)

    for pid in folder_patient_ids:
        row = [pid] + [1 if pid in sheet_reg_ids[sn] else 0 for sn in sheet_names]
        ws.append(row)

    out_wb.save(OUTPUT_XLSX)
    print(f"\n저장 완료: {OUTPUT_XLSX} ({len(folder_patient_ids)}행 x {len(sheet_names)}열)")


if __name__ == "__main__":
    main()
