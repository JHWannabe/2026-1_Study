"""
patientID_check_missing_by_sheet 폴더의 CSV 파일들을 하나의 엑셀 파일로 통합.
각 CSV 파일이 통합 엑셀의 한 시트가 된다 (시트명 = 파일명).
"""

import os
import csv
import re
import openpyxl

BASE_DIR = r"D:\데이터서비스팀 요청(이홍선)"
MISSING_DIR = os.path.join(BASE_DIR, "patientID_check_missing_by_sheet")
OUTPUT_XLSX = os.path.join(BASE_DIR, "patientID_check_missing_통합.xlsx")

INVALID_SHEETNAME_CHARS = re.compile(r'[\\/:*?\[\]]')


def sanitize_sheet_name(name):
    """엑셀 시트명 제약(특수문자 금지, 31자 제한)에 맞게 변환"""
    name = INVALID_SHEETNAME_CHARS.sub("_", name)
    return name[:31]


def main():
    csv_files = sorted(f for f in os.listdir(MISSING_DIR) if f.lower().endswith(".csv"))
    if not csv_files:
        print(f"CSV 파일이 없습니다: {MISSING_DIR}")
        return

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for csv_name in csv_files:
        sheet_name = sanitize_sheet_name(os.path.splitext(csv_name)[0])
        ws = wb.create_sheet(title=sheet_name)

        with open(os.path.join(MISSING_DIR, csv_name), newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            for row in reader:
                ws.append(row)

        print(f"[{sheet_name}] {ws.max_row - 1}건 반영")

    wb.save(OUTPUT_XLSX)
    print(f"\n통합 엑셀 저장 완료: {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
