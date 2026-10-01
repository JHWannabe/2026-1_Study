"""
Pooled dataset 20260717_VBA value.xlsx의 4개 vba_ 시트에서 source_file 컬럼 값에
실제 PNG 파일(VBA_Report_1 또는 VBA_Report_2 폴더)로 연결되는 하이퍼링크를 단다.

경로는 워크북 파일 기준 상대경로(VBA_Report_1/xxx.png 또는 VBA_Report_2/xxx.png)로 저장한다.

사용법:
    python add_vba_report_hyperlinks.py
"""

import os

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
POOLED_PATH = os.path.join(BASE_DIR, "..", "..", "data", "Pooled dataset 20260717_VBA value.xlsx")
DATA_DIR = os.path.join(BASE_DIR, "..", "..", "data")
REPORT_DIRS = ["VBA_Report_1", "VBA_Report_2"]

SHEET_NAMES = [
    "vba_anterior_height",
    "vba_middle_height",
    "vba_posterior_height",
    "vba_height_loss_ratio",
]


def build_filename_to_relpath() -> dict:
    mapping = {}
    for d in REPORT_DIRS:
        abs_dir = os.path.join(DATA_DIR, d)
        for fn in os.listdir(abs_dir):
            if not fn.lower().endswith(".png"):
                continue
            mapping[fn] = f"{d}/{fn}"
    return mapping


def main():
    filename_map = build_filename_to_relpath()
    print(f"파일명 매핑 {len(filename_map)}건 생성")

    wb = openpyxl.load_workbook(POOLED_PATH)

    for sheet_name in SHEET_NAMES:
        ws = wb[sheet_name]
        header = [c.value for c in ws[1]]
        col_idx = header.index("source_file") + 1

        linked = 0
        missing = 0
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            fn = cell.value
            if not fn:
                continue
            rel_path = filename_map.get(fn)
            if rel_path is None:
                missing += 1
                continue
            cell.hyperlink = rel_path
            cell.style = "Hyperlink"
            linked += 1

        print(f"{sheet_name}: 링크 {linked}건, 매핑 실패 {missing}건")

    wb.save(POOLED_PATH)
    wb.close()
    print("저장 완료:", POOLED_PATH)


if __name__ == "__main__":
    main()
