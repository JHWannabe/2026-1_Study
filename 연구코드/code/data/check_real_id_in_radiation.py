"""
Pooled dataset(cleaned_dataset 시트)의 Research_study_id가
강남_DLO_Results.xlsx / 신촌_DLO_Results.xlsx 파일에 존재하는지 확인하고
결과(존재 여부 + VBA_Report 이미지 링크)를 Pooled dataset 파일의 새로운 시트
('radiation_data_match')에 저장한다.

cleaned_dataset 시트의 Research_study_id 컬럼과
두 DLO_Results.xlsx 파일의 'Sheet1' 시트 'PatientID' 컬럼(4행 헤더, 5행부터 데이터)을
institution 기준으로 비교하고, 매칭되는 행의 VBA_Report 하이퍼링크(PNG 경로)를 함께 기록한다.

사용법:
    python check_real_id_in_radiation.py
"""

import os
import re

import openpyxl

POOLED_PATH = "../../data/Pooled dataset 20260717.xlsx"
RESULT_SHEET_NAME = "radiation_data_match"

INSTITUTIONS = {
    "Kangnam": {
        "dlo_xlsx": r"E:\영상제공\강남\강남_결과\강남_DLO_Results.xlsx",
        "base_dir": r"E:\영상제공\강남\강남_결과",
    },
    "Sinchon": {
        "dlo_xlsx": r"E:\영상제공\신촌\신촌_결과\신촌_DLO_Results.xlsx",
        "base_dir": r"E:\영상제공\신촌\신촌_결과",
    },
}

HYPERLINK_RE = re.compile(r'HYPERLINK\("([^"]+)"')


def load_patient_info(dlo_xlsx_path: str, base_dir: str) -> dict:
    """DLO_Results.xlsx에서 PatientID -> VBA_Report PNG 절대경로 매핑을 만든다.
    (동일 PatientID가 여러 행이면 첫 번째 행을 사용)"""
    wb = openpyxl.load_workbook(dlo_xlsx_path)
    ws = wb["Sheet1"]
    header = [c.value for c in ws[4]]
    idx_pid = header.index("PatientID")
    idx_vba = header.index("VBA_Report")

    mapping = {}
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        pid_val = row[idx_pid].value
        if pid_val is None:
            continue
        try:
            pid = int(pid_val)
        except (TypeError, ValueError):
            continue
        if pid in mapping:
            continue  # 첫 번째 행만 사용

        vba_val = row[idx_vba].value
        vba_path = None
        if isinstance(vba_val, str):
            m = HYPERLINK_RE.search(vba_val)
            if m:
                rel_path = m.group(1)
                if rel_path.startswith("./"):
                    rel_path = rel_path[2:]
                vba_path = os.path.normpath(os.path.join(base_dir, rel_path))
        mapping[pid] = vba_path

    wb.close()
    return mapping


def build_match_results(pooled_path: str, patient_info_by_institution: dict):
    """cleaned_dataset 시트를 순서대로 읽으면서 Research_study_id별 존재 여부/VBA 링크를 계산한다."""
    wb = openpyxl.load_workbook(pooled_path, read_only=True, data_only=True)
    ws = wb["cleaned_dataset"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx_study_id = header.index("Research_study_id")
    idx_institution = header.index("institution")

    results = []
    for row in rows:
        study_id = row[idx_study_id]
        if study_id is None:
            continue
        study_id = int(study_id)
        institution = row[idx_institution]

        patient_info = patient_info_by_institution.get(institution, {})
        vba_path = patient_info.get(study_id)
        exists = "Yes" if study_id in patient_info else "No"

        results.append((study_id, institution, exists, vba_path))

    wb.close()
    return results


def write_result_sheet(pooled_path: str, results: list):
    """Pooled dataset 파일에 결과를 새 시트로 추가(기존에 있으면 갱신)한다."""
    wb = openpyxl.load_workbook(pooled_path)

    if RESULT_SHEET_NAME in wb.sheetnames:
        del wb[RESULT_SHEET_NAME]
    ws = wb.create_sheet(RESULT_SHEET_NAME)

    ws.append(["Research_study_id", "institution", "in_radiation_data", "VBA_Report_link"])
    for study_id, institution, exists, vba_path in results:
        row_idx = ws.max_row + 1
        ws.append([study_id, institution, exists, vba_path])
        if vba_path and os.path.exists(vba_path):
            cell = ws.cell(row=row_idx, column=4)
            cell.hyperlink = vba_path
            cell.style = "Hyperlink"

    wb.save(pooled_path)
    wb.close()


def main():
    patient_info_by_institution = {
        inst: load_patient_info(cfg["dlo_xlsx"], cfg["base_dir"])
        for inst, cfg in INSTITUTIONS.items()
    }

    results = build_match_results(POOLED_PATH, patient_info_by_institution)

    write_result_sheet(POOLED_PATH, results)

    total = len(results)
    yes_count = sum(1 for r in results if r[2] == "Yes")
    no_count = sum(1 for r in results if r[2] == "No")
    link_count = sum(1 for r in results if r[3])

    print(f"총 {total}건 중 존재: {yes_count}건 / 미존재: {no_count}건")
    print(f"VBA_Report 링크 확보: {link_count}건")
    print(f"결과가 '{RESULT_SHEET_NAME}' 시트에 저장되었습니다: {POOLED_PATH}")


if __name__ == "__main__":
    main()
