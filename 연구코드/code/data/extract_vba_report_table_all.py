"""
VBA_Report_1, VBA_Report_2 폴더에 있는 모든 VBA report PNG(505장, 환자당 1~6장)를
전부 OCR로 추출해 Pooled dataset의 4개 vba_ 시트(anterior/middle/posterior height, HLR)를
기존 "환자당 1행"에서 "리포트(파일)당 1행"으로 완전히 교체한다.

배경:
- 기존 vba_ 시트는 DLO_Results.xlsx에서 PatientID당 첫 번째 행만 사용해 만들어졌기 때문에
  환자당 1개 리포트(phase)만 반영되어 있었다.
- 실제로는 같은 환자, 같은 촬영(Accession Number/Study Date)에서 조영 phase가 다른
  여러 장의 리포트(Plain/Arterial/Portal/Delay 등)가 존재하며, 폴더에 파일명
  "{Research_study_id}_{institution}.png", "..._2.png", "..._3.png" ... 형태로 남아있다.
- 이번 작업은 그 리포트를 전부 사용하여 각 리포트를 별도의 행으로 저장한다.

행 식별 컬럼: Real_ID, Research_study_id, institution, source_file (파일명으로 구분)

사용법:
    python extract_vba_report_table_all.py
"""

import os
import re

import openpyxl
import pytesseract
from PIL import Image

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
POOLED_PATH = os.path.join(BASE_DIR, "..", "..", "data", "Pooled dataset 20260717_VBA value.xlsx")
REPORT_DIRS = [
    os.path.join(BASE_DIR, "..", "..", "data", "VBA_Report_1"),
    os.path.join(BASE_DIR, "..", "..", "data", "VBA_Report_2"),
]

RESULT_SHEET_NAMES = {
    "Anterior": "vba_anterior_height",
    "Middle": "vba_middle_height",
    "Posterior": "vba_posterior_height",
    "HLR": "vba_height_loss_ratio",
}

LEVELS = [
    "T1", "T2", "T3", "T4", "T5", "T6", "T7", "T8", "T9", "T10",
    "T11", "T12", "L1", "L2", "L3", "L4", "L5",
]
METRICS = ["Anterior", "Middle", "Posterior", "HLR"]

# Osteo_VBA_1.png (1827x2579 고정 템플릿) 기준 표 영역 좌표
TABLE_CROP = (1000, 1100, 1827, 2350)  # left, top, right, bottom
ROW0_TOP = 51
ROW_STEP = 71.0
COL_RANGES = {
    "Anterior": (180, 300),
    "Middle": (330, 450),
    "Posterior": (480, 600),
    "HLR": (630, 760),
}

FILENAME_RE = re.compile(r"^(\d+)_([A-Za-z]+)(?:_(\d+))?\.png$")


def load_real_id_map(pooled_path: str) -> dict:
    """cleaned_dataset 시트에서 Research_study_id -> Real_ID 매핑을 만든다."""
    wb = openpyxl.load_workbook(pooled_path, read_only=True, data_only=True)
    ws = wb["cleaned_dataset"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx_study_id = header.index("Research_study_id")
    idx_real_id = header.index("Real_ID")

    mapping = {}
    for row in rows:
        study_id = row[idx_study_id]
        if study_id is None:
            continue
        mapping[int(study_id)] = row[idx_real_id]

    wb.close()
    return mapping


def list_report_files(report_dirs: list) -> list:
    """두 폴더의 모든 PNG를 (study_id, institution, suffix, abs_path, filename) 리스트로 반환."""
    files = []
    for d in report_dirs:
        for fn in sorted(os.listdir(d)):
            m = FILENAME_RE.match(fn)
            if not m:
                continue
            study_id = int(m.group(1))
            institution = m.group(2)
            suffix = int(m.group(3)) if m.group(3) else 1
            files.append((study_id, institution, suffix, os.path.join(d, fn), fn))
    return files


def extract_table_from_image(image_path: str) -> dict:
    """VBA report PNG에서 17 x 4 표를 OCR로 읽어 {level: {metric: str}} 형태로 반환."""
    im = Image.open(image_path)
    table_img = im.crop(TABLE_CROP)
    data = pytesseract.image_to_data(table_img, output_type=pytesseract.Output.DICT, config="--psm 11")

    grid = {level: {metric: None for metric in METRICS} for level in LEVELS}

    for text, top, left in zip(data["text"], data["top"], data["left"]):
        text = text.strip()
        if not text:
            continue

        row_idx = round((top - ROW0_TOP) / ROW_STEP)
        if row_idx < 0 or row_idx >= len(LEVELS):
            continue

        metric = None
        for m_name, (x0, x1) in COL_RANGES.items():
            if x0 <= left <= x1:
                metric = m_name
                break
        if metric is None:
            continue

        level = LEVELS[row_idx]
        existing = grid[level][metric]
        if existing is None or len(text) > len(existing):
            grid[level][metric] = text

    return grid


def write_result_sheets(pooled_path: str, results: list):
    """지표별로 vba_ 시트를 전부 새로 만들어 저장한다 (리포트 파일당 1행)."""
    wb = openpyxl.load_workbook(pooled_path)

    for metric, sheet_name in RESULT_SHEET_NAMES.items():
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)
        ws.append(["Real_ID", "Research_study_id", "institution", "source_file"] + LEVELS)

        for real_id, study_id, institution, source_file, grid in results:
            row_vals = [real_id, study_id, institution, source_file]
            if grid is None:
                row_vals += [None] * len(LEVELS)
            else:
                row_vals += [grid[lvl][metric] for lvl in LEVELS]
            ws.append(row_vals)

    wb.save(pooled_path)
    wb.close()


def main():
    real_id_map = load_real_id_map(POOLED_PATH)
    files = list_report_files(REPORT_DIRS)
    print(f"총 {len(files)}개 리포트 파일 발견")

    results = []
    missing_real_id = []
    ocr_failed = []

    for i, (study_id, institution, suffix, abs_path, fn) in enumerate(files, 1):
        real_id = real_id_map.get(study_id)
        if real_id is None:
            missing_real_id.append((study_id, institution, fn))

        try:
            grid = extract_table_from_image(abs_path)
        except Exception as e:
            ocr_failed.append((fn, str(e)))
            grid = None

        results.append((real_id, study_id, institution, fn, grid))

        if i % 50 == 0:
            print(f"  {i}/{len(files)} 처리 완료")

    write_result_sheets(POOLED_PATH, results)

    print(f"총 {len(results)}개 리포트 처리 완료")
    print(f"Real_ID 매핑 실패: {len(missing_real_id)}건")
    if missing_real_id:
        print("  매핑 실패 예시:", missing_real_id[:10])
    print(f"OCR 실패: {len(ocr_failed)}건")
    if ocr_failed:
        print("  OCR 실패 예시:", ocr_failed[:10])
    sheet_list = ", ".join(RESULT_SHEET_NAMES.values())
    print(f"결과가 다음 시트에 저장되었습니다 ({sheet_list}): {POOLED_PATH}")


if __name__ == "__main__":
    main()
